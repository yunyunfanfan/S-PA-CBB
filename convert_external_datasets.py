#!/usr/bin/env python3
"""Convert public fatigue databases to the common curve-level JSON schema.

The common schema stores stress amplitude (or stress range when the source only
reports range), cycles to failure, and a run-out flag.  No interpolation or
synthetic points are introduced during conversion.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _number(value, default=""):
    if isinstance(value, (int, float)):
        return value
    match = re.search(r"[-+]?\d+(?:\.\d+)?", _text(value))
    return float(match.group()) if match else default


def _write(curves, output: Path, source: str, license_name: str, notes: str):
    payload = {
        "source": source,
        "license": license_name,
        "conversion_notes": notes,
        "curves": curves,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, default=str), encoding="utf-8")
    print(json.dumps({
        "output": str(output),
        "curves": len(curves),
        "points": sum(len(curve["points"]) for curve in curves),
        "runouts": sum(point[2] for curve in curves for point in curve["points"]),
    }, ensure_ascii=False, indent=2))


def convert_hea(source: Path, output: Path):
    workbook = load_workbook(source, read_only=True, data_only=True)
    summary = workbook["HCF summary"]
    rows = summary.iter_rows(values_only=True)
    headers = [_text(value) for value in next(rows)]
    metadata = {}
    for row in rows:
        if not _text(row[0]).isdigit():
            continue
        metadata[int(row[0])] = {key: value for key, value in zip(headers, row) if key and value not in (None, "")}

    sheet = workbook["HCF individual dataset"]
    first = list(next(sheet.iter_rows(values_only=True)))
    starts = []
    for col, value in enumerate(first):
        match = re.search(r"ID\s*=\s*(\d+)", _text(value), flags=re.I)
        if match:
            starts.append((col, int(match.group(1))))
    curves = []
    for col, dataset_id in starts:
        points = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            try:
                life = float(row[col + 1])
                stress = float(row[col + 2])
            except (TypeError, ValueError, IndexError):
                continue
            if life > 0 and stress > 0:
                points.append([stress, life, 0])
        meta = metadata.get(dataset_id, {})
        composition = _text(meta.get("Composition")) or f"HEA-{dataset_id}"
        curves.append({
            "id": f"HEA2022:HCF:{dataset_id}",
            "name": composition,
            "material_family": "HEA",
            "points": points,
            "conditions": {
                "material": composition,
                "am_type": _text(meta.get("Processing history")),
                "load_ratio": _text(meta.get("R")),
                "test_type": _text(meta.get("Testing type")),
                "temperature": _text(meta.get("Tesing temperature (K)")),
            },
            "metadata": meta,
        })
    _write(
        curves, output,
        "https://doi.org/10.24435/materialscloud:s6-39",
        "Materials Cloud record terms",
        "Only the HCF stress-life worksheets are converted; LCF and FCGR records are outside the S-N task.",
    )


def convert_welded(source_dir: Path, output: Path):
    sn = json.loads((source_dir / "S-N.json").read_text())["data"]
    parameters = json.loads((source_dir / "parameter.json").read_text())["data"]
    by_id = defaultdict(list)
    for row in sn:
        try:
            dataset_id = int(row["dataset_id"])
            life = float(row["life_n"])
            stress_range = float(row["stress_range"])
            runout = int(bool(row.get("runout", 0)))
        except (KeyError, TypeError, ValueError):
            continue
        if life > 0 and stress_range > 0:
            by_id[dataset_id].append([stress_range, life, runout])
    meta_by_id = {int(row["dataset_id"]): row for row in parameters if row.get("dataset_id") is not None}
    curves = []
    for dataset_id, points in sorted(by_id.items()):
        meta = meta_by_id.get(dataset_id, {})
        material = _text(meta.get("base_material")) or "Unknown welded material"
        curves.append({
            "id": f"WELD2025:{dataset_id}",
            "name": f"{material} | {_text(meta.get('welding_joint'))}",
            "material_family": "Welded joint",
            "points": points,
            "conditions": {
                "material": material,
                "am_type": _text(meta.get("welding_method")),
                "load_ratio": _text(meta.get("load_ratio")),
                "test_type": _text(meta.get("fatigue_test_type")),
                "temperature": _text(meta.get("fatigue_temp")),
                "surface_treatment": _text(meta.get("processing")),
            },
            "metadata": meta,
        })
    _write(
        curves, output,
        "https://doi.org/10.6084/m9.figshare.29254265.v2",
        "CC BY 4.0",
        "The source reports stress range in MPa; values are retained as stress range rather than silently halved.",
    )


def convert_cma(source: Path, output: Path):
    raw = json.loads(source.read_text())
    root = raw.get("fatiguedata_cma2022", raw)
    curves = []
    dataset_id = 0
    for article_index, article in enumerate(root.get("articles", []), start=1):
        article_meta = article.get("metadata", {})
        for local_index, dataset in enumerate(article.get("scidata", {}).get("datasets", []), start=1):
            fatigue = dataset.get("fatigue", {})
            if _text(fatigue.get("fdata_type")).lower() != "sn":
                continue
            points = []
            for row in fatigue.get("fat_data", []):
                try:
                    life, stress = float(row[0]), float(row[1])
                    runout = int(bool(row[2])) if len(row) > 2 else 0
                except (TypeError, ValueError, IndexError):
                    continue
                if life > 0 and stress > 0:
                    points.append([stress, life, runout])
            dataset_id += 1
            materials = dataset.get("materials", {})
            testing = dataset.get("testing", {})
            processing = dataset.get("processing", {})
            name = _text(materials.get("mat_name")) or _text(materials.get("mat_name2")) or f"CMA-{dataset_id}"
            process_steps = [
                _text(step.get("type")) for step in processing.get("proc_para", [])
                if isinstance(step, dict) and _text(step.get("type"))
            ]
            curves.append({
                "id": f"CMA2022:{article_index}:{local_index}",
                "name": name,
                "material_family": _text(materials.get("mat_type")) or "Complex metallic alloy",
                "points": points,
                "conditions": {
                    "material": name,
                    "am_type": "; ".join(process_steps),
                    "load_ratio": _text((testing.get("fat_r") or [""])[0]),
                    "test_type": _text(testing.get("fat_type")),
                    "temperature": _text((testing.get("fat_temp") or [""])[0]),
                },
                "metadata": {
                    **article_meta,
                    "article_index": article_index,
                    "dataset_index": local_index,
                    "rating_score": dataset.get("score", []),
                    "atomic_structure": materials.get("atomic_struct", ""),
                },
            })
    _write(
        curves, output,
        "https://doi.org/10.6084/m9.figshare.23007362.v2",
        "CC BY 4.0",
        "Only S-N datasets are converted. CMA entries also present in HEA2022 are removed later by exact-curve deduplication.",
    )


def main():
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="dataset", required=True)
    hea = subparsers.add_parser("hea")
    hea.add_argument("source", type=Path)
    hea.add_argument("output", type=Path)
    welded = subparsers.add_parser("welded")
    welded.add_argument("source_dir", type=Path)
    welded.add_argument("output", type=Path)
    cma = subparsers.add_parser("cma")
    cma.add_argument("source", type=Path)
    cma.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.dataset == "hea":
        convert_hea(args.source, args.output)
    elif args.dataset == "welded":
        convert_welded(args.source_dir, args.output)
    else:
        convert_cma(args.source, args.output)


if __name__ == "__main__":
    main()
