#!/usr/bin/env python3
"""Convert the official FatigueData-AM2022 Excel workbook to pilot JSON."""

import argparse
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def text(v):
    return "" if v is None else str(v).strip()


def norm(v):
    return " ".join(text(v).lower().replace("\n", " ").split())


def convert(source: Path, output: Path):
    wb = load_workbook(source, read_only=True, data_only=True)
    sn = wb["S-N"]
    sn_rows = sn.iter_rows(values_only=True)
    sn_header = [norm(v) for v in next(sn_rows)]
    id_col = next(i for i, h in enumerate(sn_header) if "dataset id" in h)
    life_col = next(i for i, h in enumerate(sn_header) if "life" in h and "cycle" in h)
    stress_col = next(i for i, h in enumerate(sn_header) if "stress amplitude" in h)
    runout_col = next(i for i, h in enumerate(sn_header) if "runout" in h)
    points = defaultdict(list)
    for row in sn_rows:
        try:
            did = int(row[id_col])
            life, stress = float(row[life_col]), float(row[stress_col])
            runout = int(bool(row[runout_col]))
        except (TypeError, ValueError, IndexError):
            continue
        if life > 0 and stress > 0:
            points[did].append([stress, life, runout])

    para = wb["parameter"]
    rows = para.iter_rows(values_only=True)
    groups = list(next(rows))
    headers = list(next(rows))
    names = []
    current_group = ""
    for group, header in zip(groups, headers):
        if text(group):
            current_group = norm(group)
        names.append((current_group + ": " + norm(header)).strip(": "))
    metadata = {}
    for row in rows:
        try:
            did = int(row[0])
        except (TypeError, ValueError):
            continue
        metadata[did] = {name: value for name, value in zip(names, row) if value not in (None, "")}

    def find(meta, fragments):
        for key, value in meta.items():
            if all(f in key for f in fragments):
                return text(value)
        return ""

    curves = []
    for did, pts in sorted(points.items()):
        meta = metadata.get(did, {})
        material = find(meta, ("materials", "name of the material")) or find(meta, ("name of the material",))
        alloy_family = material.split("-")[0].split()[0] if material else "Unknown-AM"
        curves.append({
            "id": f"FatigueData-AM2022:{did}",
            "name": material or f"AM dataset {did}",
            "material_family": alloy_family,
            "points": pts,
            "conditions": {
                "material": material,
                "am_type": find(meta, ("am", "type")),
                "load_ratio": find(meta, ("testing", "load ratio")),
                "test_type": find(meta, ("testing", "type")),
                "temperature": find(meta, ("testing", "temperature")),
                "surface_treatment": find(meta, ("processing", "surface")),
                "heat_treatment": find(meta, ("processing", "heat")),
            },
            "metadata": meta,
        })
    payload = {
        "source": "https://doi.org/10.6084/m9.figshare.22337629.v2",
        "license": "CC BY 4.0",
        "official_file": source.name,
        "curves": curves,
    }
    output.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({
        "curves": len(curves), "points": sum(len(c["points"]) for c in curves),
        "runouts": sum(p[2] for c in curves for p in c["points"]), "output": str(output)
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("source", type=Path)
    ap.add_argument("output", type=Path)
    args = ap.parse_args()
    convert(args.source, args.output)
